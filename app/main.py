from fastapi import FastAPI, UploadFile, File, Form, HTTPException, Response
from fastapi.responses import FileResponse, JSONResponse
from pydantic import BaseModel
import openai
import os
import tempfile
import asyncio
import io
from fpdf import FPDF
from PyPDF2 import PdfReader
import docx
from dotenv import load_dotenv

# Load environment variables
load_dotenv()

# FastAPI app instance
app = FastAPI()

# Azure OpenAI API Credentials
OPENAI_API_KEY = os.getenv("AZURE_OPENAI_KEY")
OPENAI_ENDPOINT = os.getenv("AZURE_OPENAI_ENDPOINT")
MODEL_DEPLOYMENT = os.getenv("MODEL_DEPLOYMENT_NAME")
VECTOR_STORE_ID = os.getenv("VECTOR_STORE_ID")

# Configure OpenAI API
openai.api_type = "azure"
openai.api_base = OPENAI_ENDPOINT
openai.api_version = "2024-02-20"
openai.api_key = OPENAI_API_KEY

# Pydantic Model for API Input
class JobDescriptionRequest(BaseModel):
    job_description: str
    applicant_details: str

def extract_text_from_pdf(content: bytes) -> str:
    """Extracts text from a PDF file using PyPDF2."""
    reader = PdfReader(io.BytesIO(content))
    text = "\n".join([page.extract_text() or "" for page in reader.pages])
    return text.strip()

def extract_text_from_docx(content: bytes) -> str:
    """Extracts text from a Word document (.docx) using python-docx."""
    document = docx.Document(io.BytesIO(content))
    return "\n".join([para.text for para in document.paragraphs]).strip()

def extract_text_from_txt(content: bytes) -> str:
    """Decodes plain text content as UTF-8."""
    return content.decode("utf-8").strip()

async def send_to_hubspot(applicant_email: str, resume_status: str):
    """Simulates sending the resume status to HubSpot."""
    await asyncio.sleep(1)
    return {"status": "sent", "email": applicant_email, "resume_status": resume_status}

async def generate_resume(job_description: str, applicant_details: str) -> str:
    """
    Uses the fine-tuned GPT-4o model and vector store to generate a tailored resume.
    """
    try:
        # Retrieve relevant resume sections from the vector store
        search_results = openai.beta.vector_stores.files.search(
            vector_store_id=VECTOR_STORE_ID, query=job_description, max_results=5
        )
        relevant_content = "\n".join([doc.get("text", "") for doc in search_results.get("data", [])])

        # Generate a tailored resume using the fine-tuned model
        response = openai.ChatCompletion.create(
            engine=MODEL_DEPLOYMENT,
            messages=[
                {"role": "system", "content": "You generate professional, tailored resumes for job applicants."},
                {"role": "user", "content": f"Applicant Details:\n{applicant_details}\n\nJob Description:\n{job_description}\n\nRelevant Resume Sections:\n{relevant_content}"}
            ]
        )

        resume_text = response["choices"][0]["message"]["content"]
        return resume_text.strip()

    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Resume generation failed: {str(e)}")

class ResumePDF(FPDF):
    def header(self):
        """Custom header with title."""
        self.set_font("Arial", style='B', size=16)
        self.cell(0, 10, "Professional Resume", ln=True, align="C")
        self.ln(5)

    def footer(self):
        """Custom footer with page numbers."""
        self.set_y(-15)
        self.set_font("Arial", size=8)
        self.cell(0, 10, f"Page {self.page_no()}", align="C")

def generate_pdf(text: str, output_path: str):
    """Generates a professional-looking resume PDF using FPDF."""
    pdf = ResumePDF()
    pdf.add_page()
    pdf.set_font("Arial", size=12)

    sections = text.split("\n\n")  # Split text into sections for better formatting
    for section in sections:
        if section.strip():
            pdf.set_font("Arial", style='B', size=12)  # Section headers bold
            pdf.multi_cell(0, 8, section.split("\n")[0])  # First line as header
            pdf.set_font("Arial", size=11)  # Normal text
            pdf.multi_cell(0, 6, "\n".join(section.split("\n")[1:]))  # Rest of section
            pdf.ln(5)

    pdf.output(output_path)

@app.post("/generate_resume/")
async def create_resume(
    file: UploadFile = File(...),
    applicant_details: str = Form(...),
    applicant_email: str = Form(...)
):
    """
    Extracts job description from the uploaded file, generates a tailored resume using
    Azure OpenAI's fine-tuned model, and returns the resume as a PDF.
    """
    try:
        content = await file.read()
        lower_filename = file.filename.lower()

        if lower_filename.endswith(".pdf"):
            job_description = extract_text_from_pdf(content)
        elif lower_filename.endswith(".docx"):
            job_description = extract_text_from_docx(content)
        elif lower_filename.endswith(".txt"):
            job_description = extract_text_from_txt(content)
        else:
            raise HTTPException(status_code=400, detail="Unsupported file type.")

    except Exception as e:
        raise HTTPException(status_code=400, detail=f"Error reading uploaded file: {str(e)}")

    # Generate resume
    resume_text = await generate_resume(job_description, applicant_details)

    # Generate PDF
    with tempfile.NamedTemporaryFile(delete=False, suffix=".pdf") as tmp_file:
        pdf_path = tmp_file.name
    generate_pdf(resume_text, pdf_path)

    # Send update to HubSpot
    await send_to_hubspot(applicant_email, "Resume Generated")

    return FileResponse(pdf_path, media_type="application/pdf", filename="resume.pdf")

@app.get("/")
async def root():
    return JSONResponse({"message": "Resume AI Service is running"})
